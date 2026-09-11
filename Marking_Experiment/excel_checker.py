"""Reusable openpyxl checks for no-code Excel practical tasks."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict
import re
from openpyxl.utils.cell import range_boundaries

from openpyxl import load_workbook

from .checker_types import BaseChecker, CheckerResult


class ExcelChecker(BaseChecker):
    program = "excel"

    @staticmethod
    def _calculated_value(file_path: Path, sheet_name: str, cell_ref: str):
        """Read Excel's cached formula value, refreshing it with Excel when required."""
        try:
            cached = load_workbook(file_path, data_only=True, read_only=True)
            try:
                value = cached[sheet_name][cell_ref].value if sheet_name in cached.sheetnames else cached.active[cell_ref].value
            finally:
                cached.close()
            if value is not None:
                return value, "cached"
        except Exception:
            pass
        try:
            import win32com.client as win32
            excel = win32.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            workbook = excel.Workbooks.Open(str(file_path.resolve()))
            try:
                excel.CalculateFullRebuild()
                workbook.Save()
            finally:
                workbook.Close(SaveChanges=True)
                excel.Quit()
            cached = load_workbook(file_path, data_only=True, read_only=True)
            try:
                return (cached[sheet_name][cell_ref].value if sheet_name in cached.sheetnames else cached.active[cell_ref].value), "Excel recalculation"
            finally:
                cached.close()
        except Exception as exc:
            return None, f"No calculated value available: {exc}"

    def check(self, domain: str, check_type: str, target: Dict[str, Any], expected: Any, file_path: Path) -> CheckerResult:
        try:
            workbook = load_workbook(file_path, data_only=False)
        except Exception as exc:
            return CheckerResult(False, details={"reason": f"Could not read workbook: {exc}"})
        try:
            sheet_name = str((target or {}).get("sheet", ""))
            sheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
            cell_ref = str((target or {}).get("cell", ""))
            cell = sheet[cell_ref] if cell_ref else None
            wanted = str(expected or "").strip()
            actual: Any = None
            passed = False
            if check_type == "worksheet_exists":
                actual = workbook.sheetnames
                passed = wanted.casefold() in {name.casefold() for name in workbook.sheetnames}
            elif check_type == "worksheet_name":
                actual = sheet.title
                passed = actual.casefold() == wanted.casefold()
            elif check_type == "cell_value":
                actual = cell.value if cell else None
                passed = str(actual or "").strip().casefold() == wanted.casefold()
            elif check_type == "cell_contains":
                actual = cell.value if cell else None
                passed = wanted.casefold() in str(actual or "").casefold()
            elif check_type == "formula_contains":
                actual = cell.value if cell else None
                passed = isinstance(actual, str) and actual.startswith("=") and wanted.casefold() in actual.casefold()
            elif check_type == "formula_exact":
                actual = cell.value if cell else None
                passed = isinstance(actual, str) and actual.casefold().replace(" ", "") == wanted.casefold().replace(" ", "")
            elif check_type == "formula_references":
                actual = cell.value if cell else None
                references = [value.strip().casefold() for value in wanted.split(",") if value.strip()]
                passed = isinstance(actual, str) and actual.startswith("=") and all(value in actual.casefold() for value in references)
            elif check_type == "formula_function":
                actual = cell.value if cell else None
                passed = isinstance(actual, str) and re.search(rf"\b{re.escape(wanted)}\s*\(", actual, re.I) is not None
            elif check_type == "formula_operator":
                actual = cell.value if cell else None
                passed = isinstance(actual, str) and wanted in actual
            elif check_type == "formula_absolute_reference":
                actual = cell.value if cell else None
                passed = isinstance(actual, str) and wanted.upper() in actual.upper()
            elif check_type == "formula_result":
                actual, source = self._calculated_value(file_path, sheet_name, cell_ref)
                try:
                    passed = abs(float(actual) - float(expected)) < 0.000001
                except (TypeError, ValueError):
                    passed = str(actual or "").strip().casefold() == wanted.casefold()
                return CheckerResult(passed, actual, {"type": check_type, "target": target, "expected": expected, "source": source})
            elif check_type == "number_format":
                actual = cell.number_format if cell else None
                passed = wanted.casefold() in str(actual or "").casefold()
            elif check_type == "font_name":
                actual = cell.font.name if cell else None
                passed = str(actual or "").casefold() == wanted.casefold()
            elif check_type == "font_bold":
                actual = bool(cell.font.bold) if cell else False
                passed = actual == (wanted.casefold() in {"true", "yes", "1"})
            elif check_type == "font_italic":
                actual = bool(cell.font.italic) if cell else False
                passed = actual == (wanted.casefold() in {"true", "yes", "1"})
            elif check_type == "font_underline":
                actual = bool(cell.font.underline) if cell else False
                passed = actual == (wanted.casefold() in {"true", "yes", "1"})
            elif check_type == "font_color":
                color = cell.font.color if cell else None
                actual = str(getattr(color, "rgb", "") or "")
                passed = wanted.replace("#", "").casefold() in actual.casefold()
            elif check_type == "font_size":
                actual = cell.font.sz if cell else None
                passed = actual is not None and abs(float(actual) - float(expected)) < 0.01
            elif check_type == "fill_color":
                actual = cell.fill.fgColor.rgb if cell else None
                actual = str(actual or "").replace("00", "", 1).casefold()
                passed = wanted.replace("#", "").casefold() in actual
            elif check_type == "border_side":
                side, separator, style = wanted.partition(":")
                side = side.casefold() if separator else "bottom"
                style = style if separator else wanted
                border = getattr(cell.border, side, None) if cell else None
                actual = str(getattr(border, "style", "") or "")
                passed = actual.casefold() == style.casefold()
            elif check_type == "alignment":
                actual = cell.alignment.horizontal if cell else None
                passed = str(actual or "").casefold() == wanted.casefold()
            elif check_type == "vertical_alignment":
                actual = cell.alignment.vertical if cell else None
                passed = str(actual or "").casefold() == wanted.casefold()
            elif check_type == "wrap_text":
                actual = bool(cell.alignment.wrap_text) if cell else False
                passed = actual == (wanted.casefold() in {"true", "yes", "1"})
            elif check_type == "text_rotation":
                actual = cell.alignment.textRotation if cell else None
                passed = actual is not None and int(actual) == int(expected)
            elif check_type == "indent":
                actual = cell.alignment.indent if cell else None
                passed = actual is not None and abs(float(actual) - float(expected)) < 0.01
            elif check_type == "decimal_places":
                actual = cell.number_format if cell else ""
                fraction = str(actual).split(".", 1)[1].split(";", 1)[0] if "." in str(actual) else ""
                decimals = len(re.findall(r"[0#?]", fraction))
                passed = decimals == int(expected)
            elif check_type == "number_format_kind":
                actual = cell.number_format if cell else ""
                formats = str(actual).casefold()
                markers = {"currency": ("r", "$", "€", "[$"), "percentage": ("%",), "date": ("d", "m", "y")}
                passed = all(marker in formats for marker in markers.get(wanted.casefold(), (wanted.casefold(),)))
            elif check_type == "range_font_bold":
                values = [item for row in sheet[cell_ref] for item in row]
                actual = [bool(item.font.bold) for item in values]
                passed = bool(actual) and all(value == (wanted.casefold() in {"true", "yes", "1"}) for value in actual)
            elif check_type == "range_fill_color":
                values = [item for row in sheet[cell_ref] for item in row]
                actual = [str(item.fill.fgColor.rgb or "") for item in values]
                required = wanted.replace("#", "").casefold()
                passed = bool(actual) and all(required in value.casefold() for value in actual)
            elif check_type == "data_validation_range":
                validations = list(sheet.data_validations.dataValidation)
                actual = [str(item.sqref) for item in validations]
                passed = any(wanted.replace("$", "").casefold() in value.replace("$", "").casefold() for value in actual)
            elif check_type == "table_name":
                actual = list(sheet.tables)
                passed = wanted.casefold() in {name.casefold() for name in actual}
            elif check_type == "table_range":
                actual = [table.ref for table in sheet.tables.values()]
                passed = wanted.replace("$", "").casefold() in {value.replace("$", "").casefold() for value in actual}
            elif check_type == "column_width":
                column = cell.column_letter if cell else cell_ref
                actual = sheet.column_dimensions[column].width
                passed = actual is not None and float(actual) >= float(expected)
            elif check_type == "row_height":
                row = cell.row if cell else int(cell_ref)
                actual = sheet.row_dimensions[row].height
                passed = actual is not None and float(actual) >= float(expected)
            elif check_type == "freeze_panes":
                actual = str(sheet.freeze_panes or "")
                passed = actual.casefold() == wanted.casefold()
            elif check_type == "merged_range":
                actual = [str(item) for item in sheet.merged_cells.ranges]
                passed = wanted.casefold() in {item.casefold() for item in actual}
            elif check_type == "sheet_tab_color":
                color = sheet.sheet_properties.tabColor
                actual = str(getattr(color, "rgb", "") or "")
                passed = wanted.replace("#", "").casefold() in actual.casefold()
            elif check_type == "sheet_last":
                actual = workbook.sheetnames[-1]
                passed = actual.casefold() == wanted.casefold()
            elif check_type == "table_style":
                styles = [table.tableStyleInfo.name for table in sheet.tables.values() if table.tableStyleInfo]
                actual = styles
                passed = wanted.casefold() in {style.casefold() for style in styles}
            elif check_type == "sorted_range":
                # Expected: "A2:A50 ascending" or "B2:B50 descending".
                parts = wanted.split()
                range_ref = parts[0] if parts else ""
                direction = parts[1].casefold() if len(parts) > 1 else "ascending"
                cells = [item for row in sheet[range_ref] for item in row] if ":" in range_ref else []
                values = [item.value for item in cells if item.value not in (None, "")]
                comparable = [str(value).casefold() if isinstance(value, str) else value for value in values]
                actual = comparable
                passed = len(comparable) > 1 and comparable == sorted(comparable, reverse=direction == "descending")
            elif check_type == "multi_level_sort":
                # Expected: "A2:B50|ascending|ascending". The first range column is primary.
                range_ref, _, directions = wanted.partition("|")
                min_col, min_row, max_col, max_row = range_boundaries(range_ref)
                order = [part.strip().casefold() for part in directions.split("|") if part.strip()]
                rows = [tuple(sheet.cell(row, col).value for col in range(min_col, max_col + 1)) for row in range(min_row, max_row + 1)]
                rows = [row for row in rows if any(value not in (None, "") for value in row)]
                def key(row):
                    values = []
                    for value in row:
                        values.append(value.casefold() if isinstance(value, str) else value)
                    return tuple(values)
                actual = rows
                # Reverse sorting is handled per level by comparing adjacent rows.
                passed = len(rows) > 1
                for previous, current in zip(rows, rows[1:]):
                    for index, (left, right) in enumerate(zip(key(previous), key(current))):
                        if left == right:
                            continue
                        descending = index < len(order) and order[index] == "descending"
                        passed = passed and (left >= right if descending else left <= right)
                        break
                    if not passed:
                        break
            elif check_type == "conditional_formatting_present":
                actual = len(sheet.conditional_formatting)
                passed = actual > 0
            elif check_type == "conditional_formatting_range":
                actual = [str(item) for item in sheet.conditional_formatting]
                normalised_expected = wanted.replace("$", "").casefold()
                passed = any(normalised_expected in item.replace("$", "").casefold() for item in actual)
            elif check_type == "conditional_formatting_rule":
                # Expected text is matched against the rule type and formula, e.g. "cellIs|F".
                rules = []
                for item in sheet.conditional_formatting:
                    for rule in sheet.conditional_formatting[item]:
                        rules.append(f"{rule.type}|{'|'.join(str(value) for value in (rule.formula or []))}")
                actual = rules
                wanted_parts = [part.casefold() for part in wanted.split("|") if part]
                passed = any(all(part in rule.casefold() for part in wanted_parts) for rule in rules)
            elif check_type == "image_at_cell":
                anchors = []
                for image in sheet._images:
                    anchor = getattr(image, "anchor", None)
                    marker = getattr(anchor, "_from", None)
                    if marker is not None:
                        anchors.append(f"{chr(65 + marker.col)}{marker.row + 1}")
                actual = anchors
                passed = wanted.casefold() in {item.casefold() for item in anchors}
            elif check_type == "image_dimensions":
                dimensions = [(int(image.width), int(image.height)) for image in sheet._images]
                actual = dimensions
                try:
                    width, height = (int(part.strip()) for part in wanted.lower().split("x", 1))
                    passed = any(abs(current_width - width) <= 2 and abs(current_height - height) <= 2 for current_width, current_height in dimensions)
                except ValueError:
                    passed = False
            elif check_type == "chart_title":
                titles = [str(getattr(chart.title, "tx", "") or "") for chart in sheet._charts]
                actual = titles
                passed = any(wanted.casefold() in title.casefold() for title in titles)
            elif check_type == "chart_type":
                actual = [chart.__class__.__name__ for chart in sheet._charts]
                passed = any(wanted.casefold() in value.casefold() for value in actual)
            elif check_type == "chart_legend_position":
                actual = [str(getattr(getattr(chart, "legend", None), "position", "") or "") for chart in sheet._charts]
                passed = wanted.casefold() in {value.casefold() for value in actual}
            elif check_type == "chart_data_labels":
                actual = [bool(getattr(chart, "dLbls", None)) for chart in sheet._charts]
                passed = any(actual)
            elif check_type == "page_orientation":
                actual = str(sheet.page_setup.orientation or "")
                passed = actual.casefold() == wanted.casefold()
            elif check_type == "fit_to_page":
                actual = (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight)
                passed = str(actual[0]) == wanted or (wanted.casefold() == "enabled" and actual[0] is not None)
            elif check_type == "fit_to_width":
                actual = sheet.page_setup.fitToWidth
                passed = actual is not None and int(actual) == int(expected)
            elif check_type == "fit_to_height":
                actual = sheet.page_setup.fitToHeight
                passed = actual is not None and int(actual) == int(expected)
            elif check_type == "print_area":
                actual = str(sheet.print_area or "")
                passed = wanted.replace("$", "").casefold() in actual.replace("$", "").casefold()
            elif check_type == "print_title_rows":
                actual = str(sheet.print_title_rows or "")
                passed = wanted.replace("$", "").casefold() == actual.replace("$", "").casefold()
            elif check_type == "page_margins":
                side = str((target or {}).get("side", "left"))
                actual = getattr(sheet.page_margins, side, None)
                passed = actual is not None and abs(float(actual) - float(expected)) < 0.02
            elif check_type == "manual_page_break":
                actual = [item.id for item in sheet.row_breaks.brk] + [item.id for item in sheet.col_breaks.brk]
                passed = int(expected) in actual
            elif check_type == "sheet_protected":
                actual = bool(sheet.protection.sheet)
                passed = actual == (wanted.casefold() in {"true", "yes", "1"})
            elif check_type == "chart_count":
                actual = len(sheet._charts)
                passed = actual >= int(expected)
            elif check_type == "table_present":
                actual = list(sheet.tables)
                passed = bool(actual)
            return CheckerResult(passed, actual, {"type": check_type, "target": target, "expected": expected})
        finally:
            workbook.close()
